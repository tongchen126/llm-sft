import json
import matplotlib.pyplot as plt
from collections import Counter
import os
import shutil
import subprocess
import platform
import numpy as np
import seaborn as sns
from sklearn.metrics import confusion_matrix
from typing import List, Tuple, Optional
import matplotlib
import pandas as pd
import matplotlib.font_manager as fm
from openai import OpenAI
import base64
import time
import re
from pathlib import Path

def image_to_base64(image_path):
    with open(image_path, "rb") as f:
        image_base64 = base64.b64encode(f.read()).decode("utf-8")
    return image_base64


def gpt_api(
    model, system=None, user=None, image_path=None, messages=None, retry_times=5
):
    token = "irk4CnzkwB6dCF8VOOBxI2V3@2700"
    url = "http://v2.open.venus.oa.com/llmproxy"

    # 构建请求数据
    if messages is None:
        with open(image_path, "rb") as f:
            image_base64 = base64.b64encode(f.read()).decode("utf-8")

        messages = [
            {"role": "system", "content": system},
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": user},
                    {
                        "type": "image_url",
                        "image_url": {"url": f"data:image/png;base64,{image_base64}"},
                    },
                ],
            },
        ]

    client = OpenAI(base_url=url, api_key=token)

    while True:
        try:
            retry_times = retry_times - 1
            response = client.chat.completions.create(
                model=model,
                messages=messages,
            )
            break
        except Exception as e:
            pass

        if retry_times <= 0:
            raise Exception("Max retry reached...")

        time.sleep(2)

    reply = response.choices[0].message.content
    return reply

def is_font_available_insensitive(font_name):
    """
    Check if a font is available (case-insensitive).
    """
    available_fonts = [f.name.lower() for f in fm.fontManager.ttflist]
    print(sorted(available_fonts))
    return font_name.lower() in available_fonts

def install_simhei_font():
    """
    Copy SimHei.ttf font to ~/.fonts directory and refresh font cache.
    Ensures fontconfig is installed on CentOS/Ubuntu systems.
    """
    
    # Source and destination paths
    source_font = "tmp/SimHei.ttf"
    home_dir = os.path.expanduser("~")
    fonts_dir = os.path.join(home_dir, ".fonts")
    dest_font = os.path.join(fonts_dir, "SimHei.ttf")
    
    try:
        # Step 1: Check if source font exists
        if not os.path.exists(source_font):
            raise FileNotFoundError(f"Source font file not found: {source_font}")
        
        # Step 2: Detect OS and install fontconfig if needed
        print("Checking for fontconfig installation...")
        install_fontconfig()
        
        # Step 3: Create ~/.fonts directory if it doesn't exist
        if not os.path.exists(fonts_dir):
            print(f"Creating directory: {fonts_dir}")
            os.makedirs(fonts_dir, mode=0o755)
        
        # Step 4: Copy font file
        print(f"Copying {source_font} to {dest_font}")
        shutil.copy2(source_font, dest_font)

        # Step 5: Execute fc-cache -fv
        print("Refreshing font cache...")
        result = subprocess.run(
            ["fc-cache", "-fv"],
            capture_output=True,
            text=True,
            check=True
        )
        print(result.stdout)

        cache_dir = matplotlib.get_cachedir()
        print(f"Cache directory: {cache_dir}")

        # Remove the entire cache directory
        if os.path.exists(cache_dir):
            shutil.rmtree(cache_dir)
            print(f"Cache directory removed: {cache_dir}")
        else:
            print("Cache directory does not exist")

        print("✓ Font installed successfully!")
        return True

    except Exception as e:
        print(f"✗ Error: {e}")
        return False


def install_fontconfig():
    """
    Install fontconfig package on CentOS or Ubuntu if not already installed.
    """
    
    # Check if fc-cache is already available
    if shutil.which("fc-cache"):
        print("fontconfig is already installed.")
        return True
    
    print("fontconfig not found. Installing...")
    
    # Detect OS
    os_type = detect_os()
    
    try:
        if os_type == "centos":
            # CentOS/RHEL/Fedora
            print("Detected CentOS/RHEL. Installing fontconfig...")
            subprocess.run(
                ["sudo", "yum", "install", "-y", "fontconfig"],
                check=True
            )
        elif os_type == "ubuntu":
            # Ubuntu/Debian
            print("Detected Ubuntu/Debian. Installing fontconfig...")
            subprocess.run(
                ["sudo", "apt-get", "update"],
                check=True
            )
            subprocess.run(
                ["sudo", "apt-get", "install", "-y", "fontconfig"],
                check=True
            )
        else:
            print("Warning: Unknown OS. Please install fontconfig manually.")
            return False
        
        print("✓ fontconfig installed successfully!")
        return True
        
    except subprocess.CalledProcessError as e:
        print(f"✗ Failed to install fontconfig: {e}")
        return False


def detect_os():
    """
    Detect if the system is CentOS/RHEL or Ubuntu/Debian.
    """
    
    # Check /etc/os-release (modern approach)
    if os.path.exists("/etc/os-release"):
        with open("/etc/os-release", "r") as f:
            content = f.read().lower()
            if "centos" in content or "rhel" in content or "red hat" in content:
                return "centos"
            elif "ubuntu" in content or "debian" in content:
                return "ubuntu"
    
    # Check for package managers
    if shutil.which("yum") or shutil.which("dnf"):
        return "centos"
    elif shutil.which("apt-get") or shutil.which("apt"):
        return "ubuntu"
    
    return "unknown"

def plot_label_distribution(json_path, title="Label Distribution", save_path='distribution.png'):
    """
    Plot the distribution of labels showing how many labels have 1, 2, 3, etc. records.
    
    Args:
        json_path: Path to JSON file containing list of records
        title: Title for the plot (default: "Label Distribution")
    
    Returns:
        dict: Distribution data (record_count -> number_of_labels)
    """
    # Load records from JSON file
    try:
        with open(json_path, 'r', encoding='utf-8') as f:
            records = json.load(f)
    except FileNotFoundError:
        print(f"Error: File not found at {json_path}")
        return {}
    except json.JSONDecodeError:
        print(f"Error: Invalid JSON format in {json_path}")
        return {}
    
    if not records:
        print("No records to plot")
        return {}
    
    if not isinstance(records, list):
        print("Error: JSON file must contain a list of records")
        return {}
    
    # Count records per label
    label_counts = Counter(record['label'] for record in records)
    
    # Count how many labels have each record count
    # e.g., {1: 5, 2: 3, 3: 1} means 5 labels have 1 record, 3 labels have 2 records, etc.
    distribution = Counter(label_counts.values())
    
    # Sort by record count for plotting
    record_counts = sorted(distribution.keys())
    label_frequencies = [distribution[count] for count in record_counts]
    
    # Create the plot
    plt.figure(figsize=(10, 6))
    plt.bar(record_counts, label_frequencies, color='steelblue', edgecolor='black', alpha=0.7)
    
    plt.xlabel('Number of Records per Label', fontsize=12)
    plt.ylabel('Number of Labels', fontsize=12)
    plt.title(title, fontsize=14, fontweight='bold')
    plt.grid(axis='y', alpha=0.3, linestyle='--')
    
    # Add value labels on top of bars
    for x, y in zip(record_counts, label_frequencies):
        plt.text(x, y, str(y), ha='center', va='bottom', fontsize=10)
    
    # Add statistics text
    total_labels = len(label_counts)
    total_records = len(records)
    avg_records = total_records / total_labels if total_labels > 0 else 0
    
    stats_text = f'Total Labels: {total_labels}\nTotal Records: {total_records}\nAvg Records/Label: {avg_records:.2f}'
    plt.text(0.98, 0.97, stats_text, transform=plt.gca().transAxes,
            fontsize=10, verticalalignment='top', horizontalalignment='right',
            bbox=dict(boxstyle='round', facecolor='wheat', alpha=0.5))
    
    plt.tight_layout()
    plt.savefig(save_path)
    
    # Print summary
    print(f"\n{'='*50}")
    print(f"Label Distribution Summary - {json_path}")
    print(f"{'='*50}")
    print(f"Total unique labels: {total_labels}")
    print(f"Total records: {total_records}")
    print(f"Average records per label: {avg_records:.2f}")
    print(f"Distribution:")
    for count in record_counts:
        num_labels = distribution[count]
        percentage = (num_labels / total_labels) * 100
        print(f"  {num_labels} label(s) with {count} record(s) ({percentage:.1f}%)")
    print(f"{'='*50}\n")
    
    return dict(distribution)

def plot_losses_from_json(file_paths, names, output_path='loss_plot.png', 
                            x_axis='step', figsize=(10, 6)):
    """
    Extract loss values from JSON files and create a plot.
    
    Parameters:
    -----------
    file_paths : list of str
        List of paths to JSON files
    names : list of str
        List of names corresponding to each file path (for legend)
    output_path : str
        Path where the plot will be saved (default: 'loss_plot.png')
    x_axis : str
        What to use for x-axis: 'step' or 'epoch' (default: 'step')
    figsize : tuple
        Figure size (width, height) in inches
    """
    
    if len(file_paths) != len(names):
        raise ValueError("Number of file paths must match number of names")
    
    plt.figure(figsize=figsize)
    
    # Process each JSON file
    for file_path, name in zip(file_paths, names):
        # Read JSON file
        with open(file_path, 'r') as f:
            data = json.load(f)
        
        # Extract log history
        log_history = data.get('log_history', [])
        
        # Extract x-axis values and losses
        x_values = []
        losses = []
        
        for entry in log_history:
            if 'loss' in entry:  # Only include entries that have loss
                x_values.append(entry.get(x_axis, 0))
                losses.append(entry['loss'])
        
        # Plot the line
        if x_values and losses:
            plt.plot(x_values, losses, marker='o', label=name, linewidth=2, markersize=4)
        else:
            print(f"Warning: No loss data found in {file_path}")
    
    # Customize the plot
    plt.xlabel(x_axis.capitalize(), fontsize=12)
    plt.ylabel('Loss', fontsize=12)
    plt.title('Training Loss Comparison', fontsize=14, fontweight='bold')
    plt.legend(loc='best', fontsize=10)
    plt.grid(True, alpha=0.3)
    plt.tight_layout()
    
    # Save the figure
    plt.savefig(output_path, dpi=300, bbox_inches='tight')
    print(f"Plot saved to {output_path}")
    
    # Optionally display the plot
    # plt.show()
    
    plt.close()

def get_sorted_dirs(path: str, pattern: Optional[str] = None) -> List[str]:
    """
    Get sorted list of directories from a given path.
    
    Args:
        path: The directory path to search
        pattern: Optional regex pattern to filter directories (e.g., 'checkpoint.*')
    
    Returns:
        List of directory paths sorted by name
    """
    path_obj = Path(path)
    
    if not path_obj.exists() or not path_obj.is_dir():
        return []
    
    # Get all directories
    dirs = [d for d in path_obj.iterdir() if d.is_dir()]
    
    # Apply regex filter if pattern is provided
    if pattern:
        regex = re.compile(pattern)
        dirs = [d for d in dirs if regex.match(d.name)]
    
    # Sort by name
    dirs.sort(key=lambda x: x.name)
    
    return [str(d) for d in dirs]

def list_directories(path, swift=False):
    """
    List all directories in the given path.
    
    Args:
        path (str): The path to search for directories
        
    Returns:
        dict: A dictionary with directory names as keys and full paths as values
    """
    dir_dict = {}
    
    try:
        # Check if the path exists
        if not os.path.exists(path):
            raise ValueError(f"Path does not exist: {path}")
        
        # Check if the path is a directory
        if not os.path.isdir(path):
            raise ValueError(f"Path is not a directory: {path}")
        
        # List all items in the directory
        for item in sorted(os.listdir(path)):
            full_path = os.path.join(path, item)
            
            # Check if the item is a directory
            if os.path.isdir(full_path):
                abs_path = os.path.abspath(full_path)
                if swift:
                    last_checkpoint = get_sorted_dirs(get_sorted_dirs(full_path)[-1], 'checkpoint-')[-1]
                    abs_path = os.path.abspath(last_checkpoint)

                dir_dict[item] = abs_path

    except PermissionError:
        print(f"Permission denied: {path}")
    except Exception as e:
        print(f"Error: {e}")
    
    return dir_dict

def plot_prediction_heatmap(
    ground_truths: List[str],
    predictions: List[str],
    normalize: bool = False,
    figsize: tuple = None,
    cmap: str = 'Blues',
    title: str = 'Prediction vs Ground Truth',
    annot: bool = True,  # ✨ Set to False to hide numbers
    fmt: str = None,
    top_n_labels: Optional[int] = None,
    font_scale: float = 1.0,
    rotation: int = 90,
    label_mode: str = 'ground_truth_only',
    show_colorbar: bool = True,  # ✨ New: option to hide colorbar
    annot_threshold: Optional[float] = None  # ✨ New: only show numbers above threshold
) -> plt.Figure:
    """
    Create a confusion matrix heatmap.
    
    Parameters:
    -----------
    annot : bool, default=True
        If False, hide all numbers in cells (color only)
    show_colorbar : bool, default=True
        If False, hide the colorbar
    annot_threshold : float, optional
        Only annotate cells with values above this threshold
        Useful for sparse matrices
    """
    from sklearn.metrics import confusion_matrix
    
    # Get labels based on mode
    gt_labels = sorted(list(set(ground_truths)))
    pred_labels = sorted(list(set(predictions)))
    all_labels = sorted(list(set(ground_truths + predictions)))
    
    if label_mode == 'ground_truth_only':
        y_labels = gt_labels
        x_labels = all_labels
    elif label_mode == 'separate':
        y_labels = gt_labels
        x_labels = pred_labels
    elif label_mode == 'intersection':
        common = sorted(list(set(gt_labels) & set(pred_labels)))
        y_labels = x_labels = common
    else:  # union
        y_labels = x_labels = all_labels
    
    # Filter top N if specified
    if top_n_labels is not None:
        from collections import Counter
        counts = Counter(ground_truths + predictions)
        top = set([l for l, _ in counts.most_common(top_n_labels)])
        y_labels = [l for l in y_labels if l in top]
        x_labels = [l for l in x_labels if l in top]
    
    # Build confusion matrix
    if y_labels == x_labels:
        cm = confusion_matrix(ground_truths, predictions, labels=y_labels)
    else:
        cm = np.zeros((len(y_labels), len(x_labels)), dtype=int)
        y_idx = {l: i for i, l in enumerate(y_labels)}
        x_idx = {l: i for i, l in enumerate(x_labels)}
        for gt, pred in zip(ground_truths, predictions):
            if gt in y_idx and pred in x_idx:
                cm[y_idx[gt], x_idx[pred]] += 1
    
    # Normalize if requested
    if normalize:
        cm = cm.astype('float')
        row_sums = cm.sum(axis=1)[:, np.newaxis]
        row_sums[row_sums == 0] = 1
        cm = cm / row_sums
    
    # Format string
    if fmt is None:
        fmt = '.2%' if normalize else 'd'
    
    # Handle selective annotation based on threshold
    annot_data = annot
    if annot and annot_threshold is not None:
        # Create custom annotation array
        annot_data = cm.copy()
        if normalize:
            mask = cm < annot_threshold
        else:
            mask = cm < annot_threshold
        annot_data = annot_data.astype(str)
        annot_data[mask] = ''  # Hide values below threshold
    
    # Auto figure size
    if figsize is None:
        width = max(8, min(len(x_labels) * 0.6, 20))
        height = max(6, min(len(y_labels) * 0.6, 20))
        figsize = (width, height)
    
    # Create plot
    sns.set(font_scale=font_scale)
    fig, ax = plt.subplots(figsize=figsize)
    
    sns.heatmap(
        cm, 
        annot=annot_data,  # Can be True, False, or custom array
        fmt=fmt if isinstance(annot_data, bool) else '',
        cmap=cmap,
        xticklabels=x_labels, 
        yticklabels=y_labels, 
        ax=ax,
        cbar=show_colorbar,
        cbar_kws={'label': 'Percentage' if normalize else 'Count'} if show_colorbar else {},
        square=(len(y_labels) == len(x_labels)),
        linewidths=0.5,
        linecolor='gray'
    )
    
    ax.set_xlabel('Predicted Label', fontsize=12 * font_scale)
    ax.set_ylabel('True Label', fontsize=12 * font_scale)
    ax.set_title(title, fontsize=14 * font_scale, pad=20)
    
    plt.setp(ax.get_xticklabels(), rotation=rotation, ha='right', rotation_mode='anchor')
    plt.setp(ax.get_yticklabels(), rotation=0)
    
    plt.tight_layout()
    sns.reset_defaults()

    plt.rcParams["font.sans-serif"] = ["SimHei"]  # 设置字体
    plt.rcParams["axes.unicode_minus"] = False 
    
    return fig

def save_model_performance_table(model_dict, output_path='model_performance.html', 
                                    format='html', decimal_places=4):
    """
    Create and save a table of model performance metrics with best values highlighted.
    
    Parameters:
    -----------
    model_dict : dict
        Dictionary where keys are model names and values are dicts of metrics
        Example: {'model1': {'metric1': 0.95, 'metric2': 0.88}, ...}
    output_path : str
        Path where the table will be saved
    format : str
        Output format: 'html', 'excel', 'csv', or 'latex'
    decimal_places : int
        Number of decimal places to display
    
    Returns:
    --------
    df : pandas.DataFrame
        The resulting dataframe
    """
    
    # Create DataFrame from dict
    df = pd.DataFrame.from_dict(model_dict, orient='index')
    
    # Sort by model name (index)
    df = df.sort_index()
    
    # Round values
    df = df.round(decimal_places)
    
    if format == 'html':
        # Create styled HTML with highlighted best values
        def highlight_max(s):
            """Highlight the maximum in a Series yellow."""
            is_max = s == s.max()
            return ['background-color: yellow; font-weight: bold' if v else '' 
                    for v in is_max]
        
        styled_df = df.style.apply(highlight_max, axis=0)
        styled_df.to_html(output_path)
        print(f"Table saved to {output_path}")
        
    elif format == 'excel':
        from openpyxl import load_workbook
        from openpyxl.styles import PatternFill, Font
        
        # Save to Excel
        df.to_excel(output_path)
        
        # Load workbook and apply formatting
        wb = load_workbook(output_path)
        ws = wb.active
        
        # Define yellow fill and bold font
        yellow_fill = PatternFill(start_color='FFFF00', end_color='FFFF00', 
                                    fill_type='solid')
        bold_font = Font(bold=True)
        
        # Find and highlight max values in each column
        for col_idx, col in enumerate(df.columns, start=2):  # Start from 2 (B column)
            max_val = df[col].max()
            for row_idx, val in enumerate(df[col], start=2):  # Start from 2 (skip header)
                cell = ws.cell(row=row_idx, column=col_idx)
                if val == max_val:
                    cell.fill = yellow_fill
                    cell.font = bold_font
        
        wb.save(output_path)
        print(f"Table saved to {output_path}")
        
    elif format == 'csv':
        df.to_csv(output_path)
        print(f"Table saved to {output_path} (no highlighting in CSV)")
        
    elif format == 'latex':
        # For LaTeX, we'll bold the best values
        def make_bold_max(s):
            is_max = s == s.max()
            return [f'\\textbf{{{v}}}' if is_m else str(v) 
                    for v, is_m in zip(s, is_max)]
        
        df_latex = df.copy()
        for col in df_latex.columns:
            df_latex[col] = make_bold_max(df_latex[col])
        
        with open(output_path, 'w') as f:
            f.write(df_latex.to_latex(escape=False))
        print(f"Table saved to {output_path}")
    
    return df

# Set at the matplotlib level
matplotlib.rcParams['font.family'] = 'sans-serif'
matplotlib.rcParams['font.sans-serif'] = ['SimHei']
matplotlib.rcParams['axes.unicode_minus'] = False

plt.rcParams["font.sans-serif"] = ["SimHei"]  # 设置字体
plt.rcParams["axes.unicode_minus"] = False
sns.set_style("darkgrid", {
    'font.sans-serif': ['SimHei'],
    'axes.unicode_minus': False
})
if not is_font_available_insensitive("SimHei"):
    print("SimHei not available, installing...")
    install_simhei_font()
    print("Please restart program after installing the font...")
    exit(0)

# if __name__ == '__main__':
# out_path = "data/pokemon1/"
# print(plot_label_distribution(out_path + "data.json", save_path = out_path + "data.png"))
# print(plot_label_distribution(out_path + "data_eval.json", save_path = out_path + "data_eval.png"))
